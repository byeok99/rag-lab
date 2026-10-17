from __future__ import annotations

import csv, json
from pathlib import Path
from typing import Any

from .model import TestCase
from .util import safe_cell

HEADERS = ["\uae30\ub2a5\uba85 \ub300\ubd84\ub958","\uae30\ub2a5\uba85 \uc911\ubd84\ub958","\uae30\ub2a5\uba85 \uc18c\ubd84\ub958","TC_ID","\ud14c\uc2a4\ud2b8 \uc2dc\ub098\ub9ac\uc624","\uc785\ub825(\uc0ac\uc804\uc870\uac74 \ud3ec\ud568)","\uae30\ub300 \ucd9c\ub825(\uc0ac\ud6c4\uc870\uac74 \ud3ec\ud568)","\ud14c\uc2a4\ud2b8 \uacb0\uacfc(P:Pass, F:Fail)","\uc0c1\uc138 \ud14c\uc2a4\ud2b8 \uacb0\uacfc(F\uc778 \uacbd\uc6b0 \uacb0\ud568 \uc694\uc57d)","\ube44\uace0"]


def row(tc: TestCase) -> list[str]:
    input_parts=[]
    if tc.preconditions: input_parts.append("[Preconditions]\n"+tc.preconditions)
    if tc.inputs: input_parts.append("[Inputs]\n"+tc.inputs)
    input_parts.extend(f"{i}. {step}" for i,step in enumerate(tc.steps,1))
    expected=tc.expected_results or ""
    if tc.postconditions: expected += ("\n" if expected else "")+"[Postconditions]\n"+tc.postconditions
    notes=tc.notes
    if tc.requires_confirmation: notes=(notes+"\n" if notes else "")+"REQUIRES CONFIRMATION"
    return [tc.category_large,tc.category_medium,tc.category_small,tc.generated_tc_id or tc.original_tc_id,tc.scenario,"\n".join(input_parts),expected,tc.test_result if tc.origin != "generated" else "",tc.detailed_test_result if tc.origin != "generated" else "",notes]


def export_json(path: Path, cases: list[TestCase], metadata: dict[str,Any]) -> None:
    path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps({"metadata":metadata,"test_cases":[x.to_dict() for x in cases]},ensure_ascii=False,indent=2),encoding="utf-8")


def export_csv(path: Path, cases: list[TestCase]) -> None:
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f); writer.writerow(HEADERS); writer.writerows([[safe_cell(v) for v in row(tc)] for tc in cases])


def export_xlsx(path: Path, cases: list[TestCase], title: str, author: str="", reviewer: str="") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as e: raise RuntimeError("openpyxl is required: pip install -e .") from e
    wb=Workbook(); ws=wb.active; ws.title="\ud14c\uc2a4\ud2b8\ucf00\uc774\uc2a4"
    ws.merge_cells("A1:J1"); ws["A1"]=safe_cell(title); ws["A1"].font=Font(size=16,bold=True); ws["A1"].alignment=Alignment(horizontal="center")
    ws["I2"]=safe_cell(f"Author: {author}" if author else ""); ws["I3"]=safe_cell(f"Reviewer: {reviewer}" if reviewer else "")
    for col,name in enumerate(HEADERS,1): ws.cell(4,col,safe_cell(name))
    fill=PatternFill("solid",fgColor="D9EAF7"); thin=Side(style="thin",color="888888")
    for cell in ws[4]: cell.font=Font(bold=True); cell.fill=fill; cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for r,tc in enumerate(cases,5):
        for c,value in enumerate(row(tc),1):
            cell=ws.cell(r,c,safe_cell(value)); cell.alignment=Alignment(vertical="top",wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for col,width in enumerate((18,18,20,14,42,48,48,15,30,28),1): ws.column_dimensions[get_column_letter(col)].width=width
    ws.freeze_panes="E5"; ws.auto_filter.ref=f"A4:J{max(4,ws.max_row)}"; ws.sheet_view.showGridLines=False
    for col in (1,2,3):
        start=5
        while start<=ws.max_row:
            value=ws.cell(start,col).value; end=start
            while end+1<=ws.max_row and ws.cell(end+1,col).value==value: end+=1
            if value and end>start:
                ws.merge_cells(start_row=start,start_column=col,end_row=end,end_column=col); ws.cell(start,col).alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
            start=end+1
    path.parent.mkdir(parents=True,exist_ok=True); wb.save(path)

