from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .model import SourceRef, TestCase
from .util import canonical, clean_text, load_json, sha256_text, split_numbered

HEADER_ALIASES = {
    "category_large": ("\ub300\ubd84\ub958", "\uae30\ub2a5\uba85\ub300\ubd84\ub958"), "category_medium": ("\uc911\ubd84\ub958", "\uae30\ub2a5\uba85\uc911\ubd84\ub958"),
    "category_small": ("\uc18c\ubd84\ub958", "\uae30\ub2a5\uba85\uc18c\ubd84\ub958"), "original_tc_id": ("tcid", "\ud14c\uc2a4\ud2b8\ucf00\uc774\uc2a4id"),
    "scenario": ("\ud14c\uc2a4\ud2b8\uc2dc\ub098\ub9ac\uc624", "\uc2dc\ub098\ub9ac\uc624"), "raw_input_text": ("\uc785\ub825", "\uc0ac\uc804\uc870\uac74"),
    "raw_expected_output_text": ("\uae30\ub300\ucd9c\ub825", "\uc608\uc0c1\uacb0\uacfc"), "test_result": ("\ud14c\uc2a4\ud2b8\uacb0\uacfc",),
    "detailed_test_result": ("\uc0c1\uc138\ud14c\uc2a4\ud2b8\uacb0\uacfc",), "notes": ("\ube44\uace0",),
}


@dataclass
class ParseReport:
    source_document_id: str
    filename: str
    records: list[TestCase] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)


def _header_key(text: str) -> str | None:
    key = canonical(text)
    for field_name, aliases in HEADER_ALIASES.items():
        if any(canonical(a) in key for a in aliases):
            return field_name
    return None


def _build_tc(row: dict[str, Any], source: SourceRef, project: str, version: str, origin: str) -> TestCase:
    raw_in = clean_text(row.get("raw_input_text"))
    pre, steps = split_numbered(raw_in)
    expected = clean_text(row.get("raw_expected_output_text"))
    tc = TestCase(
        original_tc_id=clean_text(row.get("original_tc_id")), project_id=project,
        product_version=version, category_large=clean_text(row.get("category_large")),
        category_medium=clean_text(row.get("category_medium")), category_small=clean_text(row.get("category_small")),
        scenario=clean_text(row.get("scenario")), preconditions=pre, steps=steps,
        expected_results=expected or None, raw_input_text=raw_in, raw_expected_output_text=expected,
        source_refs=[source.__dict__], origin=origin, review_status="approved" if origin == "original" else "candidate",
        expected_result_basis="existing_tc" if expected else "unknown", requires_confirmation=not bool(expected),
        test_result=clean_text(row.get("test_result")), detailed_test_result=clean_text(row.get("detailed_test_result")),
        notes=clean_text(row.get("notes")),
    )
    source.original_tc_id = tc.original_tc_id
    tc.source_refs = [source.__dict__]
    return tc


def parse_file(path: Path, project: str, version: str, origin: str = "original") -> ParseReport:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return _parse_excel(path, project, version, origin)
    if suffix == ".csv":
        return _parse_csv(path, project, version, origin)
    if suffix == ".json":
        return _parse_json(path, project, version, origin)
    raise ValueError(f"Unsupported input: {suffix}; Markdown/TXT are extension points")


def _parse_excel(path: Path, project: str, version: str, origin: str) -> ParseReport:
    doc_id = "doc_" + sha256_text(path.read_bytes().hex())[:20]
    report = ParseReport(doc_id, path.name)
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError("openpyxl is required: pip install -e .") from e
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        sheets = []
        for ws in wb.worksheets:
            merged = {}
            for area in ws.merged_cells.ranges:
                value = ws.cell(area.min_row, area.min_col).value
                for row in range(area.min_row, area.max_row + 1):
                    for col in range(area.min_col, area.max_col + 1): merged[(row, col)] = value
            sheets.append((ws.title, ws.max_row, ws.max_column, lambda r, c, ws=ws: ws.cell(r, c).value, merged))
    else:
        try:
            import xlrd
        except ImportError as e:
            raise RuntimeError("xlrd is required for real .xls files: pip install -e .") from e
        wb = xlrd.open_workbook(path, formatting_info=True)
        sheets = []
        for ws in wb.sheets():
            merged = {}
            for rlo, rhi, clo, chi in ws.merged_cells:
                value = ws.cell_value(rlo, clo)
                for row in range(rlo + 1, rhi + 1):
                    for col in range(clo + 1, chi + 1): merged[(row, col)] = value
            sheets.append((ws.name, ws.nrows, ws.ncols, lambda r, c, ws=ws: ws.cell_value(r - 1, c - 1), merged))

    for sheet_name, max_row, max_col, value_at, merged in sheets:
        candidates = []
        for r in range(1, min(max_row, 30) + 1):
            found = {_header_key(clean_text(value_at(r, c))) for c in range(1, min(max_col, 20) + 1)}
            score = len(found - {None})
            if score >= 3 or ({"scenario", "raw_input_text", "raw_expected_output_text"} <= found): candidates.append((score, r))
        if not candidates:
            report.warnings.append({"sheet": sheet_name, "row": None, "reason": "TC header not found"})
            continue
        first_header = min(r for _, r in candidates)
        last_header = min(first_header + 2, max(r for _, r in candidates))
        mapping = {}
        for c in range(1, max_col + 1):
            joined = " ".join(clean_text(value_at(r, c)) for r in range(first_header, last_header + 1))
            key = _header_key(joined)
            if key and key not in mapping: mapping[key] = c
        # Standard A:J fallback is deliberately limited to sheets whose E:G headers were identified.
        if {"scenario", "raw_input_text", "raw_expected_output_text"} <= mapping.keys():
            for k, c in zip(("category_large","category_medium","category_small","original_tc_id","scenario","raw_input_text","raw_expected_output_text","test_result","detailed_test_result","notes"), range(1,11)):
                mapping.setdefault(k, c)
        for r in range(last_header + 1, max_row + 1):
            row = {}
            for key, c in mapping.items():
                value = value_at(r, c)
                if key in {"category_large", "category_medium", "category_small"} and (value is None or value == ""):
                    value = merged.get((r, c), "")
                row[key] = value
            if not any(clean_text(row.get(k)) for k in ("scenario", "raw_input_text", "raw_expected_output_text", "original_tc_id")):
                continue
            if not clean_text(row.get("scenario")):
                report.warnings.append({"sheet": sheet_name, "row": r, "reason": "empty scenario"})
                continue
            src = SourceRef(doc_id, path.name, sheet_name, r)
            report.records.append(_build_tc(row, src, project, version, origin))
    if not report.records:
        report.warnings.append({"sheet": "", "row": None, "reason": "no test cases parsed"})
    return report


def _map_dict(row: dict[str, Any]) -> dict[str, Any]:
    result = {}
    for key, value in row.items():
        mapped = _header_key(clean_text(key)) or key
        result[mapped] = value
    return result


def _parse_csv(path: Path, project: str, version: str, origin: str) -> ParseReport:
    raw = path.read_bytes(); doc_id = "doc_" + sha256_text(raw.hex())[:20]
    report = ParseReport(doc_id, path.name)
    text = None
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try: text = raw.decode(enc); break
        except UnicodeDecodeError: pass
    if text is None: raise ValueError("CSV encoding must be UTF-8, CP949, or EUC-KR")
    for i, row in enumerate(csv.DictReader(text.splitlines()), 2):
        mapped = _map_dict(row)
        if not clean_text(mapped.get("scenario")):
            report.warnings.append({"sheet": "CSV", "row": i, "reason": "empty scenario"}); continue
        report.records.append(_build_tc(mapped, SourceRef(doc_id, path.name, "CSV", i), project, version, origin))
    return report


def _parse_json(path: Path, project: str, version: str, origin: str) -> ParseReport:
    data = load_json(path); rows = data.get("test_cases", data) if isinstance(data, dict) else data
    if not isinstance(rows, list): raise ValueError("JSON must be an array or contain test_cases")
    doc_id = "doc_" + sha256_text(path.read_bytes().hex())[:20]; report = ParseReport(doc_id, path.name)
    for i, row in enumerate(rows, 1):
        if not isinstance(row, dict): report.warnings.append({"sheet":"JSON","row":i,"reason":"객체가 아님"}); continue
        mapped = _map_dict(row)
        if "raw_input_text" not in mapped:
            mapped["raw_input_text"] = "\n".join([clean_text(mapped.get("preconditions")), *[f"{n}. {v}" for n,v in enumerate(mapped.get("steps", []),1)]])
        if "raw_expected_output_text" not in mapped: mapped["raw_expected_output_text"] = mapped.get("expected_results")
        tc = _build_tc(mapped, SourceRef(doc_id,path.name,"JSON",i),project,version,origin)
        if mapped.get("record_id"): tc.record_id = clean_text(mapped["record_id"])
        report.records.append(tc)
    return report

